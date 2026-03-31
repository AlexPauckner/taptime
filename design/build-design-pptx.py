#!/usr/bin/env python3
"""
TapTime Design Mockup — PowerPoint/Keynote Builder
Erstellt eine editierbare PPTX mit dem Neon-Runes-Design.
Jede Kachel ist ein eigenes Objekt (verschiebbar, editierbar).
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import math

# iPhone 15 Pro Proportionen (393 x 852 pt) -> auf Slide skaliert
# Slide: 9:19.5 Ratio (iPhone), wir nutzen Custom Size
SLIDE_W = Inches(4.5)   # ~393pt skaliert
SLIDE_H = Inches(9.75)  # ~852pt skaliert

# Grid Settings
COLS = 3
ROWS = 4
GRID_PAD_X = Inches(0.3)
GRID_PAD_TOP = Inches(1.4)   # Platz fuer Header
GRID_PAD_BOTTOM = Inches(1.0)  # Platz fuer Nav
TILE_GAP = Inches(0.12)

# Berechne Tile-Groesse
grid_w = SLIDE_W - 2 * GRID_PAD_X
grid_h = SLIDE_H - GRID_PAD_TOP - GRID_PAD_BOTTOM
tile_w = (grid_w - (COLS - 1) * TILE_GAP) / COLS
tile_h = (grid_h - (ROWS - 1) * TILE_GAP) / ROWS

def hex_to_rgb(hex_str):
    hex_str = hex_str.lstrip('#')
    return RGBColor(int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16))

def darken(hex_str, factor=0.7):
    hex_str = hex_str.lstrip('#')
    r = int(int(hex_str[0:2], 16) * factor)
    g = int(int(hex_str[2:4], 16) * factor)
    b = int(int(hex_str[4:6], 16) * factor)
    return RGBColor(min(r, 255), min(g, 255), min(b, 255))

# ===========================================
# DESIGN-DATEN
# ===========================================

# Grid-Anordnung: 3x4
tiles = [
    # Zeile 1: Warm → Cool (Amber → Magenta → Deep Indigo)
    {"name": "Kita",          "emoji": "🏫", "color": "#f59e0b", "row": 0, "col": 0},
    {"name": "Brainstorm",    "emoji": "💡", "color": "#c026d3", "row": 0, "col": 1},
    {"name": "KI-Ökosystem",  "emoji": "⚙️", "color": "#4f46e5", "row": 0, "col": 2},

    # Zeile 2: Gruen → Warm (Emerald → Lime → Orange)
    {"name": "Hausarbeit",    "emoji": "🧹", "color": "#059669", "row": 1, "col": 0},
    {"name": "Kochen",        "emoji": "🍳", "color": "#65a30d", "row": 1, "col": 1},
    {"name": "Einkaufen",     "emoji": "🛒", "color": "#ea580c", "row": 1, "col": 2},

    # Zeile 3: Warm → Cool (Rose → Indigo → Violet)
    {"name": "Kinder",        "emoji": "👶", "color": "#e11d48", "row": 2, "col": 0},
    {"name": "EIINH",         "emoji": "🎬", "color": "#6366f1", "row": 2, "col": 1},
    {"name": "Bilderbuch",    "emoji": "📖", "color": "#7c3aed", "row": 2, "col": 2},

    # Zeile 4: Warm → Cool (Gold → Teal → Emerald)
    {"name": "Jasmin",        "emoji": "💛", "color": "#eab308", "row": 3, "col": 0},
    {"name": "Freizeit",      "emoji": "🎮", "color": "#0d9488", "row": 3, "col": 1},
    {"name": "Sport",         "emoji": "🏃", "color": "#16a34a", "row": 3, "col": 2},
]


def add_rounded_rect(slide, left, top, width, height, fill_color, corner_radius=Inches(0.15)):
    """Abgerundetes Rechteck mit Farbe"""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()  # Kein Rahmen

    # Corner-Radius setzen (als Prozent der kuerzeren Seite)
    # In python-pptx: shape.adjustments[0] = Wert von 0 bis 1
    # 0.15 inch bei ~1.2 inch Tile = ca. 12%
    try:
        adj = shape.adjustments
        if len(adj) > 0:
            min_dim = min(width, height)
            radius_pct = int(corner_radius / min_dim * 100000)
            adj[0] = min(radius_pct / 100000, 0.15)
    except:
        pass

    return shape


def add_text_box(slide, left, top, width, height, text, font_size, color, bold=False, alignment=PP_ALIGN.CENTER):
    """Textbox hinzufuegen"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.alignment = alignment

    return txBox


def build_presentation():
    prs = Presentation()

    # Custom Slide Size (iPhone-Proportionen)
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # Blank Layout
    blank_layout = prs.slide_layouts[6]  # Blank

    # =====================
    # SLIDE 1: Design Mockup
    # =====================
    slide = prs.slides.add_slide(blank_layout)

    # Background: Dark Navy
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(0x0d, 0x0d, 0x1a)

    # Header: "TapTime"
    add_text_box(
        slide,
        Inches(0), Inches(0.4),
        SLIDE_W, Inches(0.5),
        "TapTime",
        22, RGBColor(0xFF, 0xFF, 0xFF),
        bold=True
    )

    # Status
    add_text_box(
        slide,
        Inches(0), Inches(0.85),
        SLIDE_W, Inches(0.3),
        "Kein Timer aktiv",
        11, RGBColor(0x88, 0x88, 0x88)
    )

    # Tiles
    for tile_data in tiles:
        row = tile_data["row"]
        col = tile_data["col"]

        x = GRID_PAD_X + col * (tile_w + TILE_GAP)
        y = GRID_PAD_TOP + row * (tile_h + TILE_GAP)

        color = hex_to_rgb(tile_data["color"])

        # Tile Background (abgerundetes Rechteck)
        tile_shape = add_rounded_rect(slide, x, y, tile_w, tile_h, color)

        # Emoji (gross, zentriert oben)
        emoji_box = add_text_box(
            slide,
            x, y + Inches(0.15),
            tile_w, Inches(0.55),
            tile_data["emoji"],
            28, RGBColor(0xFF, 0xFF, 0xFF)
        )
        emoji_box.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

        # Label (klein, unten)
        # Dunkle Farben: weisse Schrift, helle Farben: dunkle Schrift
        r, g, b = color
        brightness = (r * 299 + g * 587 + b * 114) / 1000
        label_color = RGBColor(0xFF, 0xFF, 0xFF) if brightness < 160 else RGBColor(0x11, 0x11, 0x11)

        label_box = add_text_box(
            slide,
            x, y + tile_h - Inches(0.45),
            tile_w, Inches(0.35),
            tile_data["name"],
            11, label_color,
            bold=True
        )

    # Bottom Nav Bar
    nav_y = SLIDE_H - Inches(0.85)
    nav_height = Inches(0.65)

    # Nav Background
    nav_bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), nav_y,
        SLIDE_W, nav_height + Inches(0.2)
    )
    nav_bg.fill.solid()
    nav_bg.fill.fore_color.rgb = RGBColor(0x0d, 0x0d, 0x1a)
    nav_bg.line.fill.background()

    # Nav Line
    nav_line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), nav_y,
        SLIDE_W, Emu(12700)  # 1pt Linie
    )
    nav_line.fill.solid()
    nav_line.fill.fore_color.rgb = RGBColor(0x33, 0x33, 0x44)
    nav_line.line.fill.background()

    # Nav Items
    nav_items = [
        ("⏱️\nTimer", True),
        ("📊\nHeute", False),
        ("💾\nExport", False),
    ]
    nav_item_w = SLIDE_W / 3
    for i, (text, active) in enumerate(nav_items):
        color = RGBColor(0xFF, 0xFF, 0xFF) if active else RGBColor(0x66, 0x66, 0x66)
        box = add_text_box(
            slide,
            nav_item_w * i, nav_y + Inches(0.08),
            nav_item_w, nav_height,
            text,
            10, color,
            bold=active
        )

    # =====================
    # SLIDE 2: Farbpalette-Referenz
    # =====================
    slide2 = prs.slides.add_slide(blank_layout)
    bg2 = slide2.background
    fill2 = bg2.fill
    fill2.solid()
    fill2.fore_color.rgb = RGBColor(0x0d, 0x0d, 0x1a)

    add_text_box(
        slide2,
        Inches(0.3), Inches(0.3),
        SLIDE_W - Inches(0.6), Inches(0.5),
        "Farbpalette — Referenz",
        18, RGBColor(0xFF, 0xFF, 0xFF),
        bold=True,
        alignment=PP_ALIGN.LEFT
    )

    add_text_box(
        slide2,
        Inches(0.3), Inches(0.8),
        SLIDE_W - Inches(0.6), Inches(0.4),
        "Jede Zeile = harmonischer Farbverlauf links → rechts",
        12, RGBColor(0x88, 0x88, 0x88),
        alignment=PP_ALIGN.LEFT
    )

    # Farbfelder als Referenz
    row_labels = [
        "Zeile 1: Amber → Magenta → Indigo",
        "Zeile 2: Emerald → Lime → Orange",
        "Zeile 3: Rose → Indigo → Violet",
        "Zeile 4: Gold → Teal → Emerald",
    ]

    # Alternative Farbverlaeufe zum Ausprobieren
    alt_palettes = [
        # Option A: Wie oben
        [("#f59e0b", "#c026d3", "#4f46e5"),
         ("#059669", "#65a30d", "#ea580c"),
         ("#e11d48", "#6366f1", "#7c3aed"),
         ("#eab308", "#0d9488", "#16a34a")],
        # Option B: Monochromatischer pro Zeile
        [("#6366f1", "#818cf8", "#a5b4fc"),  # Indigo-Verlauf
         ("#059669", "#34d399", "#86efac"),   # Gruen-Verlauf
         ("#e11d48", "#fb7185", "#fda4af"),   # Rose-Verlauf
         ("#eab308", "#fcd34d", "#fef08a")],  # Gold-Verlauf
        # Option C: Rainbow-Flow
        [("#ef4444", "#f97316", "#eab308"),   # Rot → Orange → Gelb
         ("#22c55e", "#0d9488", "#0ea5e9"),   # Gruen → Teal → Blau
         ("#6366f1", "#8b5cf6", "#a855f7"),   # Indigo → Violet → Purple
         ("#ec4899", "#f43f5e", "#fb923c")],  # Pink → Rose → Orange
    ]

    for row_idx, label in enumerate(row_labels):
        y_pos = Inches(1.4) + row_idx * Inches(0.9)

        add_text_box(
            slide2, Inches(0.3), y_pos,
            SLIDE_W - Inches(0.6), Inches(0.25),
            label,
            10, RGBColor(0xAA, 0xAA, 0xAA),
            alignment=PP_ALIGN.LEFT
        )

        for col_idx in range(3):
            tile_data = tiles[row_idx * 3 + col_idx]
            sw = Inches(1.1)
            sx = Inches(0.3) + col_idx * (sw + Inches(0.1))
            sy = y_pos + Inches(0.28)

            shape = add_rounded_rect(slide2, sx, sy, sw, Inches(0.45), hex_to_rgb(tile_data["color"]))

            add_text_box(
                slide2, sx, sy + Inches(0.08),
                sw, Inches(0.3),
                f"{tile_data['name']}\n{tile_data['color']}",
                8, RGBColor(0xFF, 0xFF, 0xFF),
                bold=False
            )

    # Alternative Paletten
    add_text_box(
        slide2,
        Inches(0.3), Inches(5.2),
        SLIDE_W - Inches(0.6), Inches(0.5),
        "Alternative Farbverläufe",
        14, RGBColor(0xFF, 0xFF, 0xFF),
        bold=True,
        alignment=PP_ALIGN.LEFT
    )

    option_labels = ["Option B: Monochrom", "Option C: Rainbow"]
    for opt_idx, (palette, opt_label) in enumerate(zip(alt_palettes[1:], option_labels)):
        base_y = Inches(5.8) + opt_idx * Inches(1.8)

        add_text_box(
            slide2, Inches(0.3), base_y,
            SLIDE_W - Inches(0.6), Inches(0.25),
            opt_label,
            11, RGBColor(0xCC, 0xCC, 0xCC),
            bold=True,
            alignment=PP_ALIGN.LEFT
        )

        for row_idx, row_colors in enumerate(palette):
            y_pos = base_y + Inches(0.3) + row_idx * Inches(0.35)
            for col_idx, color in enumerate(row_colors):
                sw = Inches(1.1)
                sx = Inches(0.3) + col_idx * (sw + Inches(0.1))
                shape = add_rounded_rect(slide2, sx, y_pos, sw, Inches(0.28), hex_to_rgb(color))
                add_text_box(
                    slide2, sx, y_pos + Inches(0.03),
                    sw, Inches(0.22),
                    color,
                    7, RGBColor(0xFF, 0xFF, 0xFF)
                )

    # Speichern
    output_path = "/Users/alexanderpauckner/Code-Projekte/taptime/TapTime-Design-Mockup.pptx"
    prs.save(output_path)
    print(f"Gespeichert: {output_path}")


if __name__ == "__main__":
    build_presentation()
